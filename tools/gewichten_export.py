"""Export every variant without a weight to an Excel file the shop owner can fill in.

  python tools/gewichten_export.py [uitvoer.xlsx]

Default output: D:/Projects/verzendingen/producten_zonder_gewicht_<datum>.xlsx
Fill in column "Gewicht (gram)" and apply with tools/gewichten_import.py.
"""
import sys, os, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from shipauth import gql, STORE

Q = """query($c:String){ products(first:100, after:$c){ pageInfo{ hasNextPage endCursor } nodes{
  id legacyResourceId title handle productType vendor status
  cap: metafield(namespace:"custom",key:"brievenbus_capaciteit"){ value }
  variants(first:50){ nodes{ id title sku inventoryItem{ id measurement{ weight{ value unit } } } } } } } }"""

rows, c = [], None
while True:
    d = gql(Q, {"c": c})["data"]["products"]
    for p in d["nodes"]:
        for v in p["variants"]["nodes"]:
            w = v["inventoryItem"]["measurement"]["weight"]
            if w and w["value"]:
                continue
            rows.append([
                p["title"], "" if v["title"] == "Default Title" else v["title"], v["sku"] or "",
                p["productType"], p["vendor"], {"ACTIVE": "actief", "DRAFT": "concept", "ARCHIVED": "gearchiveerd", "UNLISTED": "niet in lijst"}.get(p["status"], p["status"]),
                int(p["cap"]["value"]) if p["cap"] else "", None,
                f"https://admin.shopify.com/store/{STORE.split('.')[0]}/products/{p['legacyResourceId']}",
                v["id"], v["inventoryItem"]["id"],
            ])
    if not d["pageInfo"]["hasNextPage"]:
        break
    c = d["pageInfo"]["endCursor"]

rows.sort(key=lambda r: (r[5] != "actief", r[5] == "gearchiveerd", r[3], r[4], r[0], r[1]))
out = sys.argv[1] if len(sys.argv) > 1 else f"D:/Projects/verzendingen/producten_zonder_gewicht_{datetime.date.today():%Y-%m-%d}.xlsx"

wb = Workbook()
ws = wb.active
ws.title = "Zonder gewicht"
head = ["Product", "Variant", "SKU", "Type", "Leverancier", "Status", "Brievenbus capaciteit (override)",
        "Gewicht (gram)", "Admin link", "Variant GID", "InventoryItem GID"]
ws.append(head)
for r in rows:
    ws.append(r)
bold = Font(name="Arial", bold=True)
yellow = PatternFill("solid", fgColor="FFFF00")
for cell in ws[1]:
    cell.font = bold
for row in ws.iter_rows(min_row=2):
    for cell in row:
        cell.font = Font(name="Arial")
    row[7].fill = yellow  # Gewicht (gram): in te vullen
    row[8].hyperlink = row[8].value
    row[8].font = Font(name="Arial", color="0000FF", underline="single")
widths = [48, 22, 16, 12, 22, 12, 14, 14, 60, 34, 36]
for i, w in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w
ws.freeze_panes = "A2"
ws.auto_filter.ref = ws.dimensions
ws.column_dimensions["J"].hidden = True
ws.column_dimensions["K"].hidden = True

lg = wb.create_sheet("Uitleg")
lines = [
    "Producten en varianten waarvan Shopify nog geen gewicht kent.",
    "Vul in de gele kolom 'Gewicht (gram)' het gewicht van 1 stuk in (inclusief verpakking).",
    "Laat de cel leeg als je het gewicht nog niet weet; die rij wordt dan overgeslagen.",
    "Zonder gewicht rekent de verzendmodule het product als 0 gram: het telt niet mee voor de gewichtsgrens.",
    "De kolom 'Brievenbus capaciteit (override)' is ter info: 0 = dit product gaat nooit in een brievenbuspakket.",
    "Toepassen: python tools/gewichten_import.py <dit bestand>   (in D:/Projects/warforge-verzending)",
    "De verborgen kolommen J en K zijn nodig voor het importeren; niet verwijderen.",
    f"Gemaakt op {datetime.date.today():%d-%m-%Y}. Aantal rijen: {len(rows)}.",
]
for l in lines:
    lg.append([l])
for row in lg.iter_rows():
    for cell in row:
        cell.font = Font(name="Arial")
        cell.alignment = Alignment(wrap_text=False)
lg.column_dimensions["A"].width = 120

wb.save(out)
print(len(rows), "varianten zonder gewicht ->", out)
